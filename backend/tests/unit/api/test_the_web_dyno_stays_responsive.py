"""Three ways one request could take the whole web process with it.

P2 row 14; API-04, API-09, API-11. The shape is the same in all three — work whose
size the caller chooses, done where it blocks everybody else.

**API-04 — an ordinary export blocks the event loop for ~100 seconds.** `export_xlsx`
is synchronous Python called straight from an `async def` handler, and the F-VIZ-04
formula-injection defence added `for cell in ws[ws.max_row]` after every `ws.append` —
a worksheet slice lookup per row, which turns a linear write superlinear. Measured in
the audit: `ws.append` alone 0.09 s for 3 000×20, `ws.append` plus the re-index
**5.67 s — 63×**; 10 000 rows took 102 s end to end. `ExportRequest` declares
`rows: max_length=50_000`, so the schema *invites* the payload that does it, and the
rate limit of 20/minute lets one caller hold the loop continuously. Any CPU-bound
stretch on the loop past `stale_running_heartbeat_timeout_seconds` also starves the
heartbeat — this repository has already paid for that lesson once, at `graph_build`.

**API-09 — pagination that bounds the response and nothing else.** Six endpoints
declare `limit`/`offset`, then run an unbounded `SELECT` and slice the Python list.
`GET /api/repos/{id}/docs?limit=1` loads all 763 `KnowledgeDoc` rows — generated prose
included — to return one object of five scalar fields.

**API-11 — an unbounded batch, run on the process that serves HTTP.** `note_ids` is
capped at 100 and `sql` at 50 000 characters, so the bounds were placed deliberately;
`queries` has no `max_length` at all, leaving `max_request_body_bytes` (10 MB) as the
only ceiling. And unlike every other heavy operation, the batch is started as an
in-process task even when Redis is configured — no run row, no heartbeat, no cancel.
"""

from __future__ import annotations

import ast
import inspect
import textwrap

import pytest

from app.api.routes import batch as batch_mod
from app.api.routes import visualizations as viz_mod


class TestTheExportDoesNotOwnTheLoop:
    """API-04."""

    def test_the_row_write_does_not_rescan_the_sheet(self) -> None:
        """Count the operation, do not time it.

        This repository has already recorded that a timing-ratio test for exactly
        this class of defect *passed against it* (`graph_build`, 2026-09-08), and the
        first draft here did the same: charging openpyxl's one-time warm-up to the
        small sample brought the ratio under the threshold.

        The quadratic term is `Worksheet.max_column`, which walks every cell written
        so far. `ws[ws.max_row]` reaches it through `iter_rows`, once per row.
        """
        from openpyxl.worksheet.worksheet import Worksheet

        from app.connectors.base import QueryResult
        from app.viz.export import export_xlsx

        original = Worksheet.max_column
        accesses = 0

        def _counted(self):
            nonlocal accesses
            accesses += 1
            return original.fget(self)

        rows = 200
        Worksheet.max_column = property(_counted)
        try:
            export_xlsx(
                QueryResult(
                    columns=[f"c{i}" for i in range(20)],
                    rows=[[f"v{r}-{c}" for c in range(20)] for r in range(rows)],
                    row_count=rows,
                )
            )
        finally:
            Worksheet.max_column = original

        assert accesses < rows, (
            f"{accesses} full-sheet scans for {rows} rows — one per row. "
            "`for cell in ws[ws.max_row]` after every append reaches "
            "`Worksheet.max_column`, which walks every cell written so far, so the "
            "cost of an export is quadratic in its rows: measured at 3 000x20, "
            "`ws.append` alone 0.09 s against 5.67 s for the pair, and a 10 000-row "
            "export took 102 s — on the event loop, against a schema that permits "
            "50 000 (API-04)"
        )

    def test_the_handler_does_not_run_it_on_the_loop(self) -> None:
        tree = ast.parse(textwrap.dedent(inspect.getsource(viz_mod.export_data)))
        exporters = {"export_xlsx", "export_csv", "export_json"}
        direct = [
            ast.unparse(node)
            for node in ast.walk(tree)
            if isinstance(node, ast.Call) and ast.unparse(node.func) in exporters
        ]
        assert not direct, (
            f"{direct} run synchronously inside an async handler, so the whole dyno "
            "serves nothing else for the duration — and a CPU-bound stretch longer "
            "than the heartbeat timeout also reads as a dead run (API-04)"
        )


class TestPaginationReachesTheDatabase:
    """API-09."""

    @pytest.mark.parametrize(
        ("module", "func"),
        [
            ("repos", "list_docs"),
            ("dashboards", "list_dashboards"),
            ("notes", "list_notes"),
            ("data_graph", "list_metrics"),
            ("data_graph", "list_relationships"),
        ],
    )
    def test_the_slice_is_not_done_in_python(self, module: str, func: str) -> None:
        import importlib

        mod = importlib.import_module(f"app.api.routes.{module}")
        source = inspect.getsource(getattr(mod, func))
        assert "[offset : offset + limit]" not in source, (
            f"{module}.{func} declares limit/offset and then loads the whole table: "
            "the limit bounds the response body and nothing else — not the query, not "
            "the rows deserialised, not the memory. On the one real production project "
            "`?limit=1` loads all 763 KnowledgeDoc rows, generated prose included, to "
            "return five scalar fields (API-09)"
        )

    @pytest.mark.asyncio
    async def test_the_store_returns_only_the_page_asked_for(self) -> None:
        """Rows out of a real database, not a parameter in a signature.

        The first draft checked that `get_latest_docs` *takes* limit/offset. Deleting
        the two lines that put them into the statement left the signature intact and
        the test green — a presence check passing for a parameter nobody reads, which
        is the third time this programme has learned the same lesson.
        """
        import uuid

        from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

        from app.knowledge.doc_store import DocStore
        from app.models.base import Base
        from app.models.knowledge_doc import KnowledgeDoc

        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        sm = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

        project_id = uuid.uuid4().hex
        async with sm() as session:
            for i in range(5):
                session.add(
                    KnowledgeDoc(
                        project_id=project_id,
                        doc_type="file",
                        source_path=f"app/f{i}.py",
                        content="x" * 100,
                    )
                )
            await session.commit()

            page = await DocStore().get_latest_docs(session, project_id, limit=2, offset=1)
            everything = await DocStore().get_latest_docs(session, project_id)
        await engine.dispose()

        assert len(everything) == 5, "the premise is wrong: the rows were not written"
        assert len(page) == 2, (
            f"asked for two rows and the store returned {len(page)}. The limit bounds "
            "the response body and nothing else — not the query, not the rows "
            "deserialised, not the memory. `content` is Text and holds the generated "
            "prose, so on the one real production project `?limit=1` loaded all 763 "
            "rows to return five scalar fields (API-09)"
        )

    def test_the_route_passes_the_bounds_on(self) -> None:
        from app.api.routes import repos as repos_mod

        tree = ast.parse(textwrap.dedent(inspect.getsource(repos_mod.list_docs)))
        calls = [
            node
            for node in ast.walk(tree)
            if isinstance(node, ast.Call) and "get_latest_docs" in ast.unparse(node.func)
        ]
        assert calls, "the route no longer reads the doc store; this guard is blind"
        passed = {kw.arg for call in calls for kw in call.keywords}
        assert {"limit", "offset"} <= passed, (
            f"the store can page and the route does not ask it to: {passed}"
        )


class TestABatchIsBoundedAndRunsOnTheWorker:
    """API-11."""

    def test_the_query_list_has_a_ceiling(self) -> None:
        field = batch_mod.BatchExecuteRequest.model_fields["queries"]
        limits = [m for m in field.metadata if getattr(m, "max_length", None)]
        assert limits, (
            "`note_ids` is capped at 100 and `sql` at 50 000 characters, so the bounds "
            "were placed deliberately — `queries` has none, leaving a 10 MB request "
            "body as the only ceiling. ~150 000 items is accepted, answered 202, and "
            "then executed against the customer's database with no run row, no "
            "heartbeat and no cancel route (API-11)"
        )

    def test_it_is_enqueued_when_a_worker_exists(self) -> None:
        """Reachable, and reachable on the right condition.

        The first draft accepted an `enqueue` call appearing anywhere in the handler,
        so putting `if False:` in front of it passed — code that exists and cannot run
        is the state half this row is about.
        """
        tree = ast.parse(textwrap.dedent(inspect.getsource(batch_mod.execute_batch)))
        guards = [
            ast.unparse(node.test)
            for node in ast.walk(tree)
            if isinstance(node, ast.If) and "enqueue" in ast.unparse(node.body)
        ]
        assert guards, (
            "every other heavy operation goes to the ARQ worker; the batch is started "
            "as an in-process asyncio task on the process that must stay responsive, "
            "with no run row, no heartbeat and no cancel route (API-11)"
        )
        assert any("is_arq_active" in g for g in guards), (
            f"the hand-off is guarded on something other than a worker existing: {guards}"
        )
